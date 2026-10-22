"""第一批：設定、使用者隔離（防護 6）、Excel 儲存（防護 4）、照片與檔名（防護 3）。"""
import logging
import os
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from stylemate_ai import access, excel_store, photos
from stylemate_ai.config import ConfigError, load_config
from stylemate_ai.excel_store import HEADERS, StorageError, UserStore


# ------------------------------------------------------------------ 設定

GOOD = "DISCORD_TOKEN=abc\nCHANNEL_ID=123\nCLAUDE_MODEL=haiku\nDATA_DIR=./data\n"


def write_env(tmp_path, text):
    p = tmp_path / ".env"
    p.write_text(text, encoding="utf-8")
    return p


def test_config_loads(tmp_path):
    cfg = load_config(write_env(tmp_path, GOOD))
    assert cfg.channel_id == 123 and cfg.claude_model == "haiku"
    assert cfg.data_dir == (tmp_path / "data").resolve()
    assert cfg.users_dir == cfg.data_dir / "users"


def test_config_errors(tmp_path):
    with pytest.raises(ConfigError, match="CHANNEL_ID"):
        load_config(write_env(tmp_path, "DISCORD_TOKEN=a\nCLAUDE_MODEL=h\nDATA_DIR=./d\n"))
    with pytest.raises(ConfigError, match="範例"):
        load_config(write_env(tmp_path, GOOD.replace("abc", "你的bot_token")))
    with pytest.raises(ConfigError, match="數字"):
        load_config(write_env(tmp_path, GOOD.replace("123", "abc")))
    with pytest.raises(ConfigError):
        load_config(tmp_path / "missing.env")


# ------------------------------------------------------------------ 防護 6：使用者隔離

def test_user_dir_is_numeric_and_under_users(tmp_path):
    p = access.user_dir(tmp_path, 123456789012345678)
    assert p == (tmp_path / "users" / "123456789012345678").resolve()


@pytest.mark.parametrize("bad", ["../x", "1/../2", -5, 0, 1.5, True, None, "12"])
def test_user_dir_rejects_non_ids(tmp_path, bad):
    with pytest.raises(access.AccessError):
        access.user_dir(tmp_path, bad)


def test_is_owner_only_matches_same_id():
    assert access.is_owner(1, 1)
    assert not access.is_owner(1, 2)
    assert not access.is_owner("1", 1)
    assert not access.is_owner(True, 1)


# ------------------------------------------------------------------ Excel 儲存

def new_store(tmp_path, uid=111, **kw):
    return UserStore(tmp_path, uid, retries=kw.pop("retries", 3), delay=0, **kw)


def add_garment(wb, **kw):
    rec = {"衣服ID": "K7M2", "名稱": "白襯衫", "狀態": "正常"}
    rec.update(kw)
    excel_store.append_row(wb, excel_store.WARDROBE, rec)


def test_creates_four_sheets_with_headers(tmp_path):
    s = new_store(tmp_path)
    s.transact(lambda wb: add_garment(wb))
    wb = load_workbook(s.xlsx)
    assert wb.sheetnames == ["衣櫃", "穿搭紀錄", "修改歷史", "個人設定"]
    for name, headers in HEADERS.items():
        assert [c.value for c in wb[name][1]] == headers


def test_appends_and_never_overwrites(tmp_path):
    s = new_store(tmp_path)
    for i in range(3):
        s.transact(lambda wb, i=i: add_garment(wb, 衣服ID=f"ID{i}"))
    ids = [r["衣服ID"] for r in s.read("衣櫃")]
    assert ids == ["ID0", "ID1", "ID2"]


def test_users_have_separate_files(tmp_path):
    a, b = new_store(tmp_path, 1), new_store(tmp_path, 2)
    a.transact(lambda wb: add_garment(wb, 名稱="A的衣服"))
    assert b.read("衣櫃") == [] and not b.exists()
    b.transact(lambda wb: add_garment(wb, 名稱="B的衣服"))
    assert [r["名稱"] for r in a.read("衣櫃")] == ["A的衣服"]
    assert a.xlsx != b.xlsx and a.xlsx.parent.name == "1" and b.xlsx.parent.name == "2"


@pytest.mark.parametrize("text", ['=HYPERLINK("http://evil/","x")', "+cmd|calc", "-2+3", "@SUM(A1)", "=1+1"])
def test_protection4_formula_text_stays_text(tmp_path, text):
    s = new_store(tmp_path)
    s.transact(lambda wb: add_garment(wb, 名稱=text, 備註=text))
    ws = load_workbook(s.xlsx)["衣櫃"]
    cell = ws.cell(row=2, column=2)
    assert cell.data_type == "s" and cell.value == "'" + text
    assert not str(cell.value).startswith(("=", "+", "-", "@"))
    # 讀回來時，單引號會被拿掉，使用者看到的仍是原文字
    assert s.read("衣櫃")[0]["名稱"] == text


def test_illegal_chars_and_long_text_do_not_crash(tmp_path):
    s = new_store(tmp_path)
    s.transact(lambda wb: add_garment(wb, 名稱="a\x00b\x07c", 備註="x" * 100000))
    row = s.read("衣櫃")[0]
    assert row["名稱"] == "abc" and len(row["備註"]) == 1000


def test_locked_file_warns_and_retries_then_succeeds(tmp_path, monkeypatch, caplog):
    s = new_store(tmp_path)
    real, calls = os.replace, {"n": 0}

    def flaky(src, dst):
        calls["n"] += 1
        if calls["n"] < 3:
            raise PermissionError("file is open in Excel")
        return real(src, dst)

    monkeypatch.setattr(excel_store.os, "replace", flaky)
    with caplog.at_level(logging.WARNING):
        s.transact(lambda wb: add_garment(wb))
    assert calls["n"] == 3 and len(s.read("衣櫃")) == 1
    assert sum("無法讀寫" in m for m in caplog.messages) == 2
    assert not list(s.dir.glob("*.tmp"))


def test_locked_forever_raises_storage_error(tmp_path, monkeypatch):
    s = new_store(tmp_path)

    def always(src, dst):
        raise PermissionError("locked")

    monkeypatch.setattr(excel_store.os, "replace", always)
    with pytest.raises(StorageError):
        s.transact(lambda wb: add_garment(wb))
    assert not list(s.dir.glob("*.tmp"))


def test_fn_exception_leaves_file_unchanged(tmp_path):
    s = new_store(tmp_path)
    s.transact(lambda wb: add_garment(wb, 衣服ID="A"))

    def boom(wb):
        add_garment(wb, 衣服ID="B")
        raise RuntimeError("fail midway")

    with pytest.raises(RuntimeError):
        s.transact(boom)
    assert [r["衣服ID"] for r in s.read("衣櫃")] == ["A"]


def test_corrupt_excel_gives_storage_error(tmp_path):
    s = new_store(tmp_path)
    s.dir.mkdir(parents=True)
    s.xlsx.write_bytes(b"this is not a zip")
    with pytest.raises(StorageError):
        s.transact(lambda wb: None)


def test_update_row_changes_only_given_columns(tmp_path):
    s = new_store(tmp_path)
    s.transact(lambda wb: add_garment(wb, 材質="棉"))
    row = s.read("衣櫃")[0]
    s.transact(lambda wb: excel_store.update_row(wb, "衣櫃", row["_row"], {"材質": "亞麻"}))
    after = s.read("衣櫃")[0]
    assert after["材質"] == "亞麻" and after["名稱"] == "白襯衫"


# ------------------------------------------------------------------ 防護 3：照片與檔名

JPG = b"\xff\xd8\xff\xe0" + b"0" * 20
PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 20
WHEN = datetime(2026, 10, 23, 14, 5, 33)


@pytest.mark.parametrize("attack", [
    "../../evil", "..\\..\\Windows\\evil", "C:\\x\\y", "a/b", "a:b*c?d\"e<f>g|h",
    "CON", "nul", "COM1", "  ...  ", "", "   ", "a\x00b\nc\td", "\u202etxt.exe",
    "白襯衫" * 50, "..", "....",
])
def test_sanitize_never_leaves_path_characters(attack):
    out = photos.sanitize_filename_part(attack)
    assert out
    assert not any(c in out for c in '/\\:*?"<>|')
    assert ".." not in out and out == out.strip(" .") and len(out) <= 20
    assert out.upper() not in {"CON", "NUL", "PRN", "AUX", "COM1"}
    assert all(ord(c) >= 32 for c in out)


def test_save_photo_name_format_and_collisions(tmp_path):
    d = tmp_path / "photos"
    names = [photos.save_photo(d, WHEN, "白襯衫", JPG).name for _ in range(3)]
    assert names == [
        "2026-10-23 14-05-33 白襯衫.jpg",
        "2026-10-23 14-05-33 白襯衫-2.jpg",
        "2026-10-23 14-05-33 白襯衫-3.jpg",
    ]


def test_real_extension_from_content(tmp_path):
    assert photos.save_photo(tmp_path / "p", WHEN, "杯", PNG).suffix == ".png"
    assert photos.detect_image_suffix(b"MZ\x90\x00 not an image") is None
    assert photos.detect_image_suffix(b"") is None


def test_path_traversal_name_stays_inside_photos(tmp_path):
    d = tmp_path / "photos"
    p = photos.save_photo(d, WHEN, "../../../evil", JPG)
    assert p.resolve().is_relative_to(d.resolve())
    assert [x.name for x in tmp_path.iterdir()] == ["photos"]


def test_delete_photo_only_inside_photos_dir(tmp_path):
    d = tmp_path / "photos"
    p = photos.save_photo(d, WHEN, "襯衫", JPG)
    outside = tmp_path / "secret.txt"
    outside.write_text("x")
    with pytest.raises(ValueError):
        photos.delete_photo_file(d, outside)
    assert outside.exists()
    assert photos.delete_photo_file(d, p) and not p.exists()
