import logging
import os
from datetime import datetime

import pytest
from openpyxl import load_workbook

from food_agent import storage
from food_agent.storage import HEADERS, StorageError, append_record

T = datetime(2026, 10, 20, 12, 30, 45)


def add(path, **kw):
    base = dict(taken_at=T, uploader="小明", food_desc="便當", ingredients="白飯、雞腿",
                weight_g=450, calories=800, photo_path="p.jpg", delay=0)
    base.update(kw)
    return append_record(path, **base)


def rows(path):
    return list(load_workbook(path).active.iter_rows(min_row=1, values_only=True))


def test_creates_file_with_headers(tmp_path):
    p = tmp_path / "log.xlsx"
    add(p)
    assert list(rows(p)[0]) == HEADERS


def test_appends_never_overwrites_and_accumulates_daily_total(tmp_path):
    p = tmp_path / "log.xlsx"
    assert add(p, calories=800) == 800
    assert add(p, calories=500, uploader="小華") == 1300  # 這一版不分人
    r = rows(p)
    assert len(r) == 3 and r[1][5] == 800 and r[2][5] == 500
    assert r[1][6] == 800 and r[2][6] == 1300 and r[2][1] == "小華"


def test_new_day_starts_from_zero(tmp_path):
    p = tmp_path / "log.xlsx"
    add(p, calories=800)
    nextday = datetime(2026, 10, 21, 8, 0, 0)
    assert add(p, taken_at=nextday, calories=300) == 300


def test_non_food_row_blank_numbers_total_unchanged(tmp_path):
    p = tmp_path / "log.xlsx"
    add(p, calories=800)
    total = add(p, food_desc="不是食物", ingredients="馬克杯", weight_g=None, calories=None)
    r = rows(p)[-1]
    assert total == 800 and r[2] == "不是食物" and r[3] == "馬克杯" and r[4] is None and r[5] is None
    assert r[6] == 800


@pytest.mark.parametrize("text", ['=HYPERLINK("http://evil/","x")', "+cmd|calc", "-2+3", "@SUM(A1)", "=1+1"])
def test_protection4_formula_text_stays_text(tmp_path, text):
    p = tmp_path / "log.xlsx"
    add(p, food_desc=text, ingredients=text, uploader="x" + text)
    wb = load_workbook(p)
    row = list(wb.active.iter_rows(min_row=2, max_row=2))[0]
    for cell in (row[2], row[3]):
        assert cell.data_type == "s"
        assert cell.value == "'" + text
        assert not str(cell.value).startswith(("=", "+", "-", "@"))


def test_illegal_control_chars_and_long_text_do_not_crash(tmp_path):
    p = tmp_path / "log.xlsx"
    add(p, food_desc="a\x00b\x07c", ingredients="x" * 100000)
    r = rows(p)[1]
    assert r[2] == "abc" and len(r[3]) == 1000


def test_locked_file_warns_and_retries_then_succeeds(tmp_path, monkeypatch, caplog):
    p = tmp_path / "log.xlsx"
    real = os.replace
    calls = {"n": 0}

    def flaky(src, dst):
        calls["n"] += 1
        if calls["n"] < 3:
            raise PermissionError("file is open in Excel")
        return real(src, dst)

    monkeypatch.setattr(storage.os, "replace", flaky)
    with caplog.at_level(logging.WARNING):
        add(p)
    assert calls["n"] == 3 and len(rows(p)) == 2
    assert sum("無法寫入" in m for m in caplog.messages) == 2
    assert not list(tmp_path.glob("*.tmp"))


def test_locked_forever_raises_storage_error_not_crash(tmp_path, monkeypatch):
    p = tmp_path / "log.xlsx"

    def always(src, dst):
        raise PermissionError("locked")

    monkeypatch.setattr(storage.os, "replace", always)
    with pytest.raises(StorageError):
        add(p, retries=3)
    assert not list(tmp_path.glob("*.tmp"))


def test_corrupt_excel_file_gives_storage_error(tmp_path):
    p = tmp_path / "log.xlsx"
    p.write_bytes(b"this is not a zip")
    with pytest.raises(StorageError):
        add(p)
