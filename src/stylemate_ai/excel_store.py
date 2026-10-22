"""每位使用者一個 Excel（wardrobe.xlsx），四張工作表。

本檔含一段防護：
  防護 4：以 = + - @ 開頭的文字當純文字（放在 _set_cell() 裡）
框外的 cell.value = value 是「沒有防護」時的預設行為，所以把整段註解掉，程式仍能執行。

寫入規則：先寫暫存檔再換掉正式檔（原子替換）；Excel 被開著（Windows 會鎖檔）時，
印警告並重試，不讓程式當掉；每位使用者一把鎖，同一個人的兩個請求不會同時改同一個檔。
"""
from __future__ import annotations

import logging
import os
import threading
import time
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from . import access

log = logging.getLogger(__name__)

WARDROBE = "衣櫃"
OUTFITS = "穿搭紀錄"
HISTORY = "修改歷史"
SETTINGS = "個人設定"

HEADERS: dict[str, list[str]] = {
    WARDROBE: [
        "衣服ID", "名稱", "種類", "顏色", "材質", "風格", "圖案", "適合季節",
        "保暖程度", "正式程度", "AI信心", "備註", "品牌", "照片位置",
        "入庫時間", "狀態", "最後修改時間",
    ],
    OUTFITS: ["紀錄ID", "日期", "記錄時間", "衣服ID清單", "場合", "天氣", "來源", "狀態", "最後修改時間"],
    HISTORY: ["歷史ID", "對象", "對象ID", "修改時間", "修改者ID", "動作", "原因", "變更欄位", "修改前", "修改後"],
    SETTINGS: ["項目", "值"],
}
MAX_CELL_CHARS = 1000
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


class StorageError(Exception):
    """Excel 一直寫不進去，或檔案本身有問題。"""


_locks: dict[int, threading.Lock] = {}
_locks_guard = threading.Lock()


def _lock_for(user_id: int) -> threading.Lock:
    with _locks_guard:
        return _locks.setdefault(user_id, threading.Lock())


# ---------------------------------------------------------------- 儲存格

def _clean(value: Any) -> Any:
    """去掉 Excel 不接受的控制字元、限制長度（避免寫入時出錯）。"""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL_CHARS]
    return value


def _set_cell(cell, raw: Any) -> None:
    value = _clean(raw)
    cell.value = value  # 沒有防護時：字串若以 = 開頭，Excel 會當成公式

    # ==== 防護 4：以 = + - @ 開頭的文字當純文字　開始 ====
    # 文字格強制設為「文字」型態，開頭是公式符號的再加一個單引號，打開 Excel 也不會被執行。
    if isinstance(value, str):
        if value.startswith(FORMULA_PREFIXES):
            cell.value = "'" + value
        cell.data_type = "s"
    # ==== 防護 4 結束 ====


def _unquote(value: Any) -> Any:
    """讀回來時，把防護 4 加的單引號拿掉。"""
    if isinstance(value, str) and value.startswith("'") and value[1:2] in FORMULA_PREFIXES and value[1:2] != "":
        return value[1:]
    return value


# ---------------------------------------------------------------- 工作簿操作（在 transact 的 fn 裡使用）

def _new_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for name, headers in HEADERS.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
    return wb


def _ensure_sheets(wb: Workbook) -> None:
    for name, headers in HEADERS.items():
        if name not in wb.sheetnames:
            wb.create_sheet(name).append(headers)


def sheet_rows(wb: Workbook, name: str) -> list[dict]:
    ws = wb[name]
    headers = [c.value for c in ws[1]]
    rows = []
    for index, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None for v in values):
            continue
        record = {h: _unquote(v) for h, v in zip(headers, values)}
        record["_row"] = index
        rows.append(record)
    return rows


def append_row(wb: Workbook, name: str, record: dict) -> None:
    """永遠加在最後一列，從不覆蓋既有的列。"""
    ws = wb[name]
    row_number = ws.max_row + 1
    for col, header in enumerate(HEADERS[name], start=1):
        _set_cell(ws.cell(row=row_number, column=col), record.get(header))


def update_row(wb: Workbook, name: str, row_number: int, updates: dict) -> None:
    ws = wb[name]
    for col, header in enumerate(HEADERS[name], start=1):
        if header in updates:
            _set_cell(ws.cell(row=row_number, column=col), updates[header])


# ---------------------------------------------------------------- 一位使用者的儲存

class UserStore:
    """某一位使用者的資料夾與 Excel。所有讀寫都只碰這位使用者自己的資料夾。"""

    def __init__(self, data_dir: Path, user_id: int, *, retries: int = 10, delay: float = 3.0) -> None:
        self.user_id = user_id
        self.dir = access.user_dir(data_dir, user_id)
        self.xlsx = self.dir / "wardrobe.xlsx"
        self.photos_dir = self.dir / "photos"
        self.retries = retries
        self.delay = delay

    def exists(self) -> bool:
        return self.xlsx.exists()

    def _retry(self, action: Callable[[], Any]) -> Any:
        for attempt in range(1, self.retries + 1):
            try:
                return action()
            except OSError as exc:  # PermissionError 等：多半是 Excel 正開著這個檔
                log.warning(
                    "⚠️ 無法讀寫使用者 %s 的 Excel（第 %d/%d 次）：%s。如果 Excel 正開著它，請先關閉；%s 秒後重試。",
                    self.user_id, attempt, self.retries, exc, self.delay,
                )
                if attempt < self.retries:
                    time.sleep(self.delay)
            except StorageError:
                raise
        raise StorageError(f"試了 {self.retries} 次仍無法讀寫使用者 {self.user_id} 的 Excel")

    def _open(self) -> Workbook:
        try:
            if self.xlsx.exists():
                wb = load_workbook(self.xlsx)
                _ensure_sheets(wb)
                return wb
            return _new_workbook()
        except OSError:
            raise
        except Exception as exc:  # 例如檔案壞了：重試也沒用
            raise StorageError(f"Excel 檔案有問題：{exc}") from exc

    def read(self, sheet: str) -> list[dict]:
        if not self.exists():
            return []
        with _lock_for(self.user_id):
            return self._retry(lambda: sheet_rows(self._open(), sheet))

    def transact(self, fn: Callable[[Workbook], Any]) -> Any:
        """在鎖裡面：打開（或建立）→ 執行 fn → 原子存檔。fn 丟出的例外會原樣往外丟、不存檔。"""
        with _lock_for(self.user_id):
            def attempt():
                wb = self._open()
                result = fn(wb)
                self.dir.mkdir(parents=True, exist_ok=True)
                tmp = self.xlsx.with_name(self.xlsx.name + ".tmp")
                try:
                    wb.save(tmp)
                    os.replace(tmp, self.xlsx)  # 要嘛整個換好，要嘛完全沒動
                except BaseException:
                    tmp.unlink(missing_ok=True)
                    raise
                return result

            return self._retry(attempt)
