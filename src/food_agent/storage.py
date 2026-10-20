"""Excel 紀錄：新的一列永遠加在最後，不覆蓋；檔案被開著時重試。

本檔含一段防護：
  防護 4：以 = + - @ 開頭的文字當純文字（放在 _write_row() 裡）
框外的 cell.value = value 是「沒有防護時」的預設行為，所以把整段註解掉，程式仍能執行。
"""
from __future__ import annotations

import logging
import os
import time
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

log = logging.getLogger(__name__)

SHEET_NAME = "紀錄"
HEADERS = [
    "日期時間", "上傳者", "食物敘述", "成份", "重量(公克)",
    "熱量(大卡)", "當日已攝取熱量(大卡)", "照片存放位置",
]
MAX_CELL_CHARS = 1000
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
CALORIES_COL = 6


class StorageError(Exception):
    """Excel 一直寫不進去。"""


def _clean(value: object) -> object:
    """去掉 Excel 不接受的控制字元、限制長度（避免寫入時出錯）。"""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:MAX_CELL_CHARS]
    return value


def _write_row(ws, row_number: int, values: list[object]) -> None:
    for col, raw in enumerate(values, start=1):
        value = _clean(raw)
        cell = ws.cell(row=row_number, column=col)
        cell.value = value  # 沒有防護時：字串若以 = 開頭，Excel 會當成公式

        # ==== 防護 4：以 = + - @ 開頭的文字當純文字　開始 ====
        # 文字格強制設為「文字」型態，開頭是公式符號的再加一個單引號，打開 Excel 也不會被執行。
        if isinstance(value, str):
            if value.startswith(FORMULA_PREFIXES):
                cell.value = "'" + value
            cell.data_type = "s"
        # ==== 防護 4 結束 ====


def _open_or_create(xlsx_path: Path):
    if xlsx_path.exists():
        wb = load_workbook(xlsx_path)
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
        return wb, ws
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADERS)
    for width, letter in zip((20, 14, 18, 30, 12, 12, 22, 60), "ABCDEFGH"):
        ws.column_dimensions[letter].width = width
    return wb, ws


def _daily_total_before(ws, day: str) -> float:
    total = 0.0
    for row in ws.iter_rows(min_row=2, values_only=True):
        when, calories = row[0], row[CALORIES_COL - 1]
        if isinstance(when, str) and when.startswith(day) and isinstance(calories, (int, float)):
            total += calories
    return total


def _append_once(xlsx_path: Path, values_builder) -> float:
    wb, ws = _open_or_create(xlsx_path)
    total, values = values_builder(ws)
    _write_row(ws, ws.max_row + 1, values)
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = xlsx_path.with_name(xlsx_path.name + ".tmp")
    try:
        wb.save(tmp)
        os.replace(tmp, xlsx_path)  # 要嘛整個換好，要嘛完全沒動
    except BaseException:
        tmp.unlink(missing_ok=True)
        raise
    return total


def append_record(
    xlsx_path: Path,
    *,
    taken_at: datetime,
    uploader: str,
    food_desc: str,
    ingredients: str,
    weight_g: float | None,
    calories: float | None,
    photo_path: str,
    retries: int = 10,
    delay: float = 3.0,
) -> float:
    """在 Excel 最後加一列，回傳「當天到這一餐為止」的已攝取熱量。
    檔案被開著（Windows 會鎖住）時，印警告並重試；一直不行才丟 StorageError。"""
    day = taken_at.strftime("%Y-%m-%d")

    def build(ws):
        total = _daily_total_before(ws, day) + (calories or 0)
        return total, [
            taken_at.strftime("%Y-%m-%d %H:%M:%S"),
            uploader,
            food_desc,
            ingredients,
            weight_g,
            calories,
            round(total, 1),
            photo_path,
        ]

    for attempt in range(1, retries + 1):
        try:
            return _append_once(xlsx_path, build)
        except OSError as exc:  # PermissionError 等：多半是 Excel 正開著這個檔
            log.warning(
                "⚠️ 無法寫入 %s（第 %d/%d 次）：%s。如果 Excel 正開著它，請先關閉；%s 秒後重試。",
                xlsx_path, attempt, retries, exc, delay,
            )
            if attempt < retries:
                time.sleep(delay)
        except Exception as exc:  # 例如 Excel 檔本身壞了：重試也沒用
            raise StorageError(f"Excel 檔案有問題：{exc}") from exc
    raise StorageError(f"試了 {retries} 次仍無法寫入 {xlsx_path}")
